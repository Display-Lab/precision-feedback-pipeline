from __future__ import annotations

import argparse
from copy import copy
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple
import sys
import uuid

import pandas as pd
from openpyxl import load_workbook


# -----------------------------------------------------------------------------
# Measure and conversion constants
# -----------------------------------------------------------------------------
TARGET_MEASURES = ["attendance-01"]
VALID_ROLES = {"faculty": "Faculty", "staff": "Staff", "student": "Student"}
VALID_ATTENDANCE = {"in_person", "remote", "not_attended"}
ATTENDED_VALUES = {"in_person", "remote"}

OVERALL_ROLE_CODE = "DLHS member"
OVERALL_SUBJECT_IDENTIFIER = 9999
GROUP_SUBJECT = "DLHS"

ROLE_GROUPS = [
    (OVERALL_ROLE_CODE, None, OVERALL_SUBJECT_IDENTIFIER),
    ("Faculty", {"Faculty"}, 1119),
    ("Staff", {"Staff"}, 1121),
    ("Student", {"Student"}, 1120),
]

ATTENDANCE_COLUMN_CANDIDATES = [
    "Attendance (In-person or remote, or not_attended)",
    "Attendance (In-person or remote)",
    "Attendance",
]

DEFAULT_MEASURE_NAMES: Dict[str, str] = {
    "attendance-01": "Overall Attendance",
}

COMPARATOR_RATES = {
    "attendance-01": 0.80,
}


# -----------------------------------------------------------------------------
# Helpers
# -----------------------------------------------------------------------------
def normalize_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    return str(value).strip()


def normalize_key(value: object) -> str:
    return normalize_text(value).lower()


def normalize_attendance_value(value: object) -> str:
    v = normalize_key(value)
    v = v.replace(" ", "_").replace("-", "_")
    return v


def normalize_role_value(value: object) -> str:
    v = normalize_key(value)
    if v in VALID_ROLES:
        return VALID_ROLES[v]
    return ""


def resolve_excel_path(path_value: str, label: str) -> Path:
    path = Path(path_value).expanduser()
    if not path.exists():
        raise FileNotFoundError(f"{label} workbook not found: {path}")
    if path.suffix.lower() not in {".xlsx", ".xlsm", ".xls"}:
        raise ValueError(f"{label} must be an Excel workbook path (.xlsx, .xlsm, .xls): {path}")
    return path


def trim_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    out.columns = [normalize_text(c) for c in out.columns]
    for col in out.columns:
        if out[col].dtype == object:
            out[col] = out[col].map(lambda x: normalize_text(x) if pd.notna(x) else "")
    return out


def build_column_lookup(columns: Sequence[str]) -> Dict[str, str]:
    return {normalize_key(c): c for c in columns}


def require_column(df: pd.DataFrame, column_name: str, sheet_label: str) -> str:
    lookup = build_column_lookup(df.columns)
    key = normalize_key(column_name)
    if key not in lookup:
        raise ValueError(
            f"Required column '{column_name}' not found in sheet '{sheet_label}'. "
            f"Found columns: {list(df.columns)}"
        )
    return lookup[key]


def require_first_matching_column(df: pd.DataFrame, candidates: Sequence[str], sheet_label: str) -> str:
    lookup = build_column_lookup(df.columns)
    for name in candidates:
        key = normalize_key(name)
        if key in lookup:
            return lookup[key]
    raise ValueError(
        f"Required column not found in sheet '{sheet_label}'. "
        f"Accepted names: {list(candidates)}. Found columns: {list(df.columns)}"
    )


def find_sheet_by_trimmed_name(sheet_names: Sequence[str], expected_name: str) -> str:
    expected = normalize_key(expected_name)
    for s in sheet_names:
        if normalize_key(s) == expected:
            return s
    raise ValueError(
        f"Required sheet '{expected_name}' is missing. "
        f"Found sheets: {list(sheet_names)}"
    )


def safe_date(value: object) -> Optional[date]:
    if value is None or normalize_text(value) == "":
        return None
    ts = pd.to_datetime(value, errors="coerce")
    if pd.isna(ts):
        return None
    return ts.date()


def list_sheet_headers(path: Path) -> Dict[str, List[str]]:
    xl = pd.ExcelFile(path)
    headers: Dict[str, List[str]] = {}
    for sheet_name in xl.sheet_names:
        preview = pd.read_excel(path, sheet_name=sheet_name, nrows=0)
        headers[sheet_name] = [normalize_text(c) for c in preview.columns.tolist()]
    return headers


def get_openpyxl_headers(ws) -> List[str]:
    return [normalize_text(ws.cell(row=1, column=c).value) for c in range(1, ws.max_column + 1)]


def header_index_map(headers: Sequence[str]) -> Dict[str, int]:
    return {normalize_key(h): i for i, h in enumerate(headers)}


def require_sheet_columns(sheet_name: str, headers: Sequence[str], required: Sequence[str]) -> Dict[str, int]:
    idx = header_index_map(headers)
    out: Dict[str, int] = {}
    missing: List[str] = []
    for name in required:
        key = normalize_key(name)
        if key not in idx:
            missing.append(name)
        else:
            out[name] = idx[key]
    if missing:
        raise ValueError(
            f"Required columns missing in sheet '{sheet_name}': {missing}. "
            f"Found headers: {list(headers)}"
        )
    return out


def row_has_any_value(row_values: Sequence[object]) -> bool:
    return any(normalize_text(v) != "" for v in row_values)


def extract_records(ws, headers: Sequence[str]) -> List[Dict[str, object]]:
    records: List[Dict[str, object]] = []
    for r in range(2, ws.max_row + 1):
        vals = [ws.cell(row=r, column=c).value for c in range(1, len(headers) + 1)]
        if not row_has_any_value(vals):
            continue
        rec = {headers[c - 1]: vals[c - 1] for c in range(1, len(headers) + 1)}
        records.append(rec)
    return records


def clear_data_rows(ws) -> None:
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)


def copy_row_style(ws, src_row: int, dst_row: int, max_col: int) -> None:
    for c in range(1, max_col + 1):
        src = ws.cell(row=src_row, column=c)
        dst = ws.cell(row=dst_row, column=c)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format is not None:
            dst.number_format = src.number_format
        if src.protection is not None:
            dst.protection = copy(src.protection)
        if src.alignment is not None:
            dst.alignment = copy(src.alignment)
        if src.fill is not None:
            dst.fill = copy(src.fill)
        if src.font is not None:
            dst.font = copy(src.font)
        if src.border is not None:
            dst.border = copy(src.border)


def write_records(ws, headers: Sequence[str], records: Sequence[Dict[str, object]], style_template_row: Optional[int]) -> None:
    max_col = len(headers)
    for idx, rec in enumerate(records, start=2):
        if style_template_row is not None:
            copy_row_style(ws, style_template_row, idx, max_col)
        for c in range(1, max_col + 1):
            header = headers[c - 1]
            ws.cell(row=idx, column=c, value=rec.get(header, None))


# -----------------------------------------------------------------------------
# Data models
# -----------------------------------------------------------------------------
@dataclass
class MeasureMeta:
    identifier: str
    name: str
    title: str
    description: str


@dataclass
class PeriodResult:
    measure: str
    practitioner_role_code: str
    subject_identifier: int
    period_start: date
    period_end: date
    numerator: int
    denominator: int

    @property
    def rate(self) -> float:
        return self.numerator / self.denominator


# -----------------------------------------------------------------------------
# Load and validate source data
# -----------------------------------------------------------------------------
def load_registry_data(path: Path) -> Tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, List[str], Dict[str, str]]:
    warnings: List[str] = []
    mapping_info: Dict[str, str] = {}
    xl = pd.ExcelFile(path)
    trimmed_sheet_map = {normalize_key(s): s for s in xl.sheet_names}

    required_sheets = ["Meetings", "People", "Attendance Events"]
    missing = [s for s in required_sheets if normalize_key(s) not in trimmed_sheet_map]
    if missing:
        raise ValueError(f"Registry workbook is missing required sheet(s): {missing}. Found: {xl.sheet_names}")

    meetings = trim_dataframe(pd.read_excel(path, sheet_name=trimmed_sheet_map[normalize_key("Meetings")], dtype=object))
    people = trim_dataframe(pd.read_excel(path, sheet_name=trimmed_sheet_map[normalize_key("People")], dtype=object))
    attendance = trim_dataframe(
        pd.read_excel(path, sheet_name=trimmed_sheet_map[normalize_key("Attendance Events")], dtype=object)
    )

    mapping_info["registry.Meetings.sheet"] = trimmed_sheet_map[normalize_key("Meetings")]
    mapping_info["registry.People.sheet"] = trimmed_sheet_map[normalize_key("People")]
    mapping_info["registry.AttendanceEvents.sheet"] = trimmed_sheet_map[normalize_key("Attendance Events")]

    m_id = require_column(meetings, "ID", "Meetings")
    m_location = require_column(meetings, "Location", "Meetings")
    m_date = require_column(meetings, "Date", "Meetings")

    p_id = require_column(people, "ID", "People")
    p_role = require_column(people, "Role", "People")
    _p_enroll = require_column(people, "LHS 701 Enrollment", "People")

    a_person = require_column(attendance, "Person ID", "Attendance Events")
    a_meeting = require_column(attendance, "Meeting ID", "Attendance Events")
    a_attendance = require_first_matching_column(attendance, ATTENDANCE_COLUMN_CANDIDATES, "Attendance Events")

    mapping_info["registry.Meetings.ID"] = m_id
    mapping_info["registry.Meetings.Location"] = m_location
    mapping_info["registry.Meetings.Date"] = m_date
    mapping_info["registry.People.ID"] = p_id
    mapping_info["registry.People.Role"] = p_role
    mapping_info["registry.AttendanceEvents.PersonID"] = a_person
    mapping_info["registry.AttendanceEvents.MeetingID"] = a_meeting
    mapping_info["registry.AttendanceEvents.Attendance"] = a_attendance

    meetings = meetings[[m_id, m_location, m_date]].copy()
    meetings.columns = ["ID", "Location", "Date"]

    people = people[[p_id, p_role]].copy()
    people.columns = ["ID", "Role"]

    attendance = attendance[[a_person, a_meeting, a_attendance]].copy()
    attendance.columns = ["Person ID", "Meeting ID", "Attendance"]

    people["ID"] = people["ID"].map(normalize_text)
    people["Role_raw"] = people["Role"].map(normalize_text)
    people["Role"] = people["Role"].map(normalize_role_value)

    invalid_roles = sorted({r for r in people["Role_raw"].tolist() if normalize_text(r) and normalize_role_value(r) == ""})
    if invalid_roles:
        warnings.append(f"Invalid role values in People: {invalid_roles}")

    role_conflicts = (
        people[people["ID"] != ""]
        .groupby("ID")["Role"]
        .nunique(dropna=False)
        .reset_index(name="n_roles")
    )
    conflicting_ids = role_conflicts[role_conflicts["n_roles"] > 1]["ID"].tolist()
    if conflicting_ids:
        warnings.append(f"Duplicate person IDs with conflicting roles found; keeping first occurrence: {conflicting_ids}")

    duplicate_people_ids = people[people["ID"] != ""]["ID"].duplicated(keep=False)
    if duplicate_people_ids.any():
        dup_people = sorted(people.loc[duplicate_people_ids, "ID"].unique().tolist())
        warnings.append(f"Duplicate People IDs found; keeping first occurrence per ID: {dup_people}")

    people = people.drop_duplicates(subset=["ID"], keep="first")

    meetings["ID"] = meetings["ID"].map(normalize_text)
    meetings["Date_parsed"] = meetings["Date"].map(safe_date)

    bad_dates = meetings[(meetings["ID"] != "") & (meetings["Date_parsed"].isna())]["ID"].tolist()
    if bad_dates:
        warnings.append(f"Meetings with invalid or missing Date will be ignored in reports: {bad_dates}")

    attendance["Person ID"] = attendance["Person ID"].map(normalize_text)
    attendance["Meeting ID"] = attendance["Meeting ID"].map(normalize_text)
    attendance["Attendance_raw"] = attendance["Attendance"].map(normalize_text)
    attendance["Attendance"] = attendance["Attendance"].map(normalize_attendance_value)

    invalid_attendance_values = sorted(
        {
            raw
            for raw, norm in zip(attendance["Attendance_raw"].tolist(), attendance["Attendance"].tolist())
            if normalize_text(raw) != "" and norm not in VALID_ATTENDANCE
        }
    )
    if invalid_attendance_values:
        warnings.append(f"Invalid attendance values in Attendance Events: {invalid_attendance_values}")

    people_ids = set(people["ID"].dropna().astype(str).map(normalize_text).tolist())
    meeting_ids = set(meetings["ID"].dropna().astype(str).map(normalize_text).tolist())

    invalid_person_ids = sorted(
        {
            pid
            for pid in attendance["Person ID"].tolist()
            if normalize_text(pid) != "" and normalize_text(pid) not in people_ids
        }
    )
    if invalid_person_ids:
        warnings.append(f"Attendance rows with Person ID not found in People: {invalid_person_ids}")

    invalid_meeting_ids = sorted(
        {
            mid
            for mid in attendance["Meeting ID"].tolist()
            if normalize_text(mid) != "" and normalize_text(mid) not in meeting_ids
        }
    )
    if invalid_meeting_ids:
        warnings.append(f"Attendance rows with Meeting ID not found in Meetings: {invalid_meeting_ids}")

    duplicate_meeting_ids = meetings[meetings["ID"] != ""]["ID"].duplicated(keep=False)
    if duplicate_meeting_ids.any():
        dup_meetings = sorted(meetings.loc[duplicate_meeting_ids, "ID"].unique().tolist())
        warnings.append(f"Duplicate Meetings IDs found; keeping all rows but date joins may be ambiguous: {dup_meetings}")

    pair_series = attendance[["Person ID", "Meeting ID"]].copy()
    pair_series["Person ID"] = pair_series["Person ID"].map(normalize_text)
    pair_series["Meeting ID"] = pair_series["Meeting ID"].map(normalize_text)
    valid_pairs = pair_series[(pair_series["Person ID"] != "") & (pair_series["Meeting ID"] != "")]
    duplicate_pairs = valid_pairs.duplicated(subset=["Person ID", "Meeting ID"], keep=False)
    if duplicate_pairs.any():
        dup_count = int(duplicate_pairs.sum())
        warnings.append(
            "Duplicate attendance records found for the same Person ID + Meeting ID pair; "
            f"all rows are retained ({dup_count} duplicate rows)."
        )

    return meetings, people, attendance, warnings, mapping_info


def default_measure_meta() -> Dict[str, MeasureMeta]:
    out: Dict[str, MeasureMeta] = {}
    for measure_id in TARGET_MEASURES:
        name = DEFAULT_MEASURE_NAMES[measure_id]
        out[measure_id] = MeasureMeta(
            identifier=measure_id,
            name=name,
            title=name,
            description="",
        )
    return out


def load_measure_reference(path: Optional[Path]) -> Tuple[Dict[str, MeasureMeta], str]:
    if path is None:
        return default_measure_meta(), "No measure reference workbook provided; using built-in measure names."

    xl = pd.ExcelFile(path)

    id_candidates = [
        "measure.identifier",
        "identifier",
        "id",
        "measure id",
        "measure_id",
    ]
    name_candidates = ["measure.name", "measure name", "name"]
    title_candidates = ["measure.title", "measure name", "title"]
    description_keywords = ["description", "logic", "rationale", "notes"]

    found: Dict[str, MeasureMeta] = {}

    for sheet in xl.sheet_names:
        df = trim_dataframe(pd.read_excel(path, sheet_name=sheet, dtype=object))
        col_lookup = build_column_lookup(df.columns)

        id_col = None
        for c in id_candidates:
            if c in col_lookup:
                id_col = col_lookup[c]
                break
        if id_col is None:
            continue

        df["__id_norm__"] = df[id_col].map(normalize_key)

        for measure_id in TARGET_MEASURES:
            if measure_id in found:
                continue

            matches = df[df["__id_norm__"] == measure_id]
            if matches.empty:
                continue

            row = matches.iloc[0]

            name_value = ""
            for c in name_candidates:
                if c in col_lookup:
                    name_value = normalize_text(row[col_lookup[c]])
                    if name_value:
                        break

            title_value = ""
            for c in title_candidates:
                if c in col_lookup:
                    title_value = normalize_text(row[col_lookup[c]])
                    if title_value:
                        break

            description_value = ""
            for col in df.columns:
                key = normalize_key(col)
                if any(k in key for k in description_keywords):
                    description_value = normalize_text(row[col])
                    if description_value:
                        break

            if not name_value:
                name_value = measure_id
            if not title_value:
                title_value = measure_id

            found[measure_id] = MeasureMeta(
                identifier=measure_id,
                name=name_value,
                title=title_value,
                description=description_value,
            )

        if len(found) == len(TARGET_MEASURES):
            break

    missing = [m for m in TARGET_MEASURES if m not in found]
    if missing:
        raise ValueError(
            "Could not find all required measure definitions in the measure reference workbook. "
            f"Missing IDs: {missing}. Ensure identifiers exist and can be matched after trimming/lowercasing."
        )

    return found, "Loaded measure metadata from reference workbook."


# -----------------------------------------------------------------------------
# Calculations
# -----------------------------------------------------------------------------
def calculate_period_results(
    meetings: pd.DataFrame,
    people: pd.DataFrame,
    attendance: pd.DataFrame,
) -> Tuple[List[PeriodResult], List[str], Dict[str, int]]:
    warnings: List[str] = []
    processing_stats: Dict[str, int] = {}

    people_clean = people[["ID", "Role"]].copy()
    meetings_clean = meetings[["ID", "Date_parsed"]].copy()

    joined = attendance.merge(
        people_clean,
        how="left",
        left_on="Person ID",
        right_on="ID",
        suffixes=("", "_person"),
    )
    joined = joined.merge(
        meetings_clean,
        how="left",
        left_on="Meeting ID",
        right_on="ID",
        suffixes=("", "_meeting"),
    )

    joined = joined.rename(columns={"Role": "Role_norm", "Date_parsed": "MeetingDate"})

    total_rows = int(len(joined))

    valid_rows = joined[
        joined["Person ID"].map(normalize_text).ne("")
        & joined["Meeting ID"].map(normalize_text).ne("")
        & joined["Attendance"].isin(VALID_ATTENDANCE)
        & joined["Role_norm"].isin(VALID_ROLES.values())
        & joined["MeetingDate"].notna()
    ].copy()

    processing_stats["attendance_rows_total"] = total_rows
    processing_stats["attendance_rows_valid"] = int(len(valid_rows))
    processing_stats["attendance_rows_unprocessed"] = int(total_rows - len(valid_rows))

    if valid_rows.empty:
        warnings.append("No valid attendance rows remain after validation filters.")
        return [], warnings, processing_stats

    results: List[PeriodResult] = []

    def aggregate_for_group(practitioner_role_code: str, roles: Optional[set], subject_identifier: int) -> None:
        if roles is None:
            filtered = valid_rows.copy()
        else:
            filtered = valid_rows[valid_rows["Role_norm"].isin(roles)].copy()

        if filtered.empty:
            return

        grouped = filtered.groupby("MeetingDate", dropna=True)
        for meeting_date, grp in grouped:
            denominator = int(len(grp))
            if denominator == 0:
                continue
            numerator = int(grp["Attendance"].isin(ATTENDED_VALUES).sum())
            results.append(
                PeriodResult(
                    measure="attendance-01",
                    practitioner_role_code=practitioner_role_code,
                    subject_identifier=subject_identifier,
                    period_start=meeting_date,
                    period_end=meeting_date,
                    numerator=numerator,
                    denominator=denominator,
                )
            )

    for role_code, role_filter, subject_id in ROLE_GROUPS:
        aggregate_for_group(role_code, role_filter, subject_id)

    role_order = {role: i for i, (role, _roles, _sid) in enumerate(ROLE_GROUPS)}
    results.sort(key=lambda x: (x.period_start.isoformat(), role_order.get(x.practitioner_role_code, 99)))
    return results, warnings, processing_stats


# -----------------------------------------------------------------------------
# Workbook writing
# -----------------------------------------------------------------------------
def build_practitioner_role_records(headers: Sequence[str]) -> List[Dict[str, object]]:
    required = [
        "PractitionerRole.identifier",
        "PractitionerRole.practitioner",
        "PractitionerRole.code",
        "PractitionerRole.organization",
    ]
    require_sheet_columns("PractitionerRole", headers, required)

    identifier_col = next(h for h in headers if normalize_key(h) == normalize_key("PractitionerRole.identifier"))
    practitioner_col = next(h for h in headers if normalize_key(h) == normalize_key("PractitionerRole.practitioner"))
    role_col = next(h for h in headers if normalize_key(h) == normalize_key("PractitionerRole.code"))
    org_col = next(h for h in headers if normalize_key(h) == normalize_key("PractitionerRole.organization"))

    records: List[Dict[str, object]] = []
    for role_code, _roles, subject_id in ROLE_GROUPS:
        rec = {h: None for h in headers}
        rec[identifier_col] = subject_id
        # The revised template leaves this field blank for role groups.
        rec[practitioner_col] = None
        rec[role_col] = role_code
        rec[org_col] = GROUP_SUBJECT
        records.append(rec)

    return records


def build_measure_records(
    headers: Sequence[str],
    measure_meta: Dict[str, MeasureMeta],
) -> List[Dict[str, object]]:
    required = [
        "Measure.name",
        "Measure.title",
        "Measure.identifier",
        "Measure.group.type",
        "Measure.improvementNotation",
    ]
    require_sheet_columns("Measure", headers, required)

    name_col = next(h for h in headers if normalize_key(h) == normalize_key("Measure.name"))
    title_col = next(h for h in headers if normalize_key(h) == normalize_key("Measure.title"))
    id_col = next(h for h in headers if normalize_key(h) == normalize_key("Measure.identifier"))
    group_type_col = next(h for h in headers if normalize_key(h) == normalize_key("Measure.group.type"))
    improve_col = next(h for h in headers if normalize_key(h) == normalize_key("Measure.improvementNotation"))

    records: List[Dict[str, object]] = []
    for measure_id in TARGET_MEASURES:
        meta = measure_meta[measure_id]
        rec = {h: None for h in headers}
        rec[name_col] = meta.name
        rec[title_col] = meta.title
        rec[id_col] = measure_id
        rec[group_type_col] = "process"
        rec[improve_col] = "increasing"
        records.append(rec)

    return records


def build_performance_records(
    headers: Sequence[str],
    existing_records: Sequence[Dict[str, object]],
    period_results: Sequence[PeriodResult],
) -> List[Dict[str, object]]:
    required = [
        "identifier",
        "measure",
        "subject",
        "period.start",
        "period.end",
        "measureScore.rate",
        "measureScore.denominator",
        "measureScore.range",
    ]
    require_sheet_columns("PerformanceMeasureReport", headers, required)

    measure_col = next(h for h in headers if normalize_key(h) == normalize_key("measure"))

    preserved = [
        rec
        for rec in existing_records
        if normalize_key(rec.get(measure_col, "")) not in set(TARGET_MEASURES)
    ]

    identifier_col = next(h for h in headers if normalize_key(h) == normalize_key("identifier"))
    subject_col = next(h for h in headers if normalize_key(h) == normalize_key("subject"))
    period_start_col = next(h for h in headers if normalize_key(h) == normalize_key("period.start"))
    period_end_col = next(h for h in headers if normalize_key(h) == normalize_key("period.end"))
    rate_col = next(h for h in headers if normalize_key(h) == normalize_key("measureScore.rate"))
    denom_col = next(h for h in headers if normalize_key(h) == normalize_key("measureScore.denominator"))
    range_col = next(h for h in headers if normalize_key(h) == normalize_key("measureScore.range"))

    new_rows: List[Dict[str, object]] = []
    for result in period_results:
        if result.denominator == 0:
            continue

        rec = {h: None for h in headers}
        rec[identifier_col] = f"pmr-{result.measure}-{result.period_start.strftime('%Y%m%d')}-{uuid.uuid4().hex[:8]}"
        rec[measure_col] = result.measure
        rec[subject_col] = int(result.subject_identifier)
        rec[period_start_col] = result.period_start
        rec[period_end_col] = result.period_end
        rec[rate_col] = float(result.rate)
        rec[denom_col] = int(result.denominator)
        rec[range_col] = "NULL"
        new_rows.append(rec)

    return preserved + new_rows


def build_comparator_records(
    headers: Sequence[str],
    existing_records: Sequence[Dict[str, object]],
    period_results: Sequence[PeriodResult],
) -> List[Dict[str, object]]:
    required = [
        "identifier",
        "measure",
        "group.subject",
        "period.start",
        "period.end",
        "measureScore.rate",
        "measureScore.denominator",
        "group.code",
        "PractitionerRole.code",
    ]
    require_sheet_columns("ComparatorMeasureReport", headers, required)

    measure_col = next(h for h in headers if normalize_key(h) == normalize_key("measure"))

    preserved = [
        rec
        for rec in existing_records
        if normalize_key(rec.get(measure_col, "")) not in set(TARGET_MEASURES)
    ]

    identifier_col = next(h for h in headers if normalize_key(h) == normalize_key("identifier"))
    group_subject_col = next(h for h in headers if normalize_key(h) == normalize_key("group.subject"))
    period_start_col = next(h for h in headers if normalize_key(h) == normalize_key("period.start"))
    period_end_col = next(h for h in headers if normalize_key(h) == normalize_key("period.end"))
    rate_col = next(h for h in headers if normalize_key(h) == normalize_key("measureScore.rate"))
    denom_col = next(h for h in headers if normalize_key(h) == normalize_key("measureScore.denominator"))
    group_code_col = next(h for h in headers if normalize_key(h) == normalize_key("group.code"))
    role_code_col = next(h for h in headers if normalize_key(h) == normalize_key("PractitionerRole.code"))

    new_rows: List[Dict[str, object]] = []
    for result in period_results:
        if result.denominator == 0:
            continue

        measure_id = result.measure
        rec = {h: None for h in headers}
        rec[identifier_col] = f"cmr-{measure_id}-{result.period_start.strftime('%Y%m%d')}-{uuid.uuid4().hex[:8]}"
        rec[measure_col] = measure_id
        rec[group_subject_col] = GROUP_SUBJECT
        rec[period_start_col] = result.period_start
        rec[period_end_col] = result.period_end
        rec[rate_col] = float(COMPARATOR_RATES[measure_id])
        rec[denom_col] = int(result.denominator)
        rec[group_code_col] = "Goal Comparator"
        rec[role_code_col] = result.practitioner_role_code
        new_rows.append(rec)

    return preserved + new_rows


def write_output_workbook(
    scaffold_template_path: Path,
    output_path: Path,
    measure_meta: Dict[str, MeasureMeta],
    period_results: Sequence[PeriodResult],
) -> Tuple[int, int, int, int]:
    wb = load_workbook(scaffold_template_path)

    practitioner_sheet_name = find_sheet_by_trimmed_name(wb.sheetnames, "PractitionerRole")
    measure_sheet_name = find_sheet_by_trimmed_name(wb.sheetnames, "Measure")
    performance_sheet_name = find_sheet_by_trimmed_name(wb.sheetnames, "PerformanceMeasureReport")
    comparator_sheet_name = find_sheet_by_trimmed_name(wb.sheetnames, "ComparatorMeasureReport")

    ws_pr = wb[practitioner_sheet_name]
    ws_meas = wb[measure_sheet_name]
    ws_perf = wb[performance_sheet_name]
    ws_comp = wb[comparator_sheet_name]

    pr_headers = get_openpyxl_headers(ws_pr)
    meas_headers = get_openpyxl_headers(ws_meas)
    perf_headers = get_openpyxl_headers(ws_perf)
    comp_headers = get_openpyxl_headers(ws_comp)

    pr_style_row = 2 if ws_pr.max_row >= 2 else None
    meas_style_row = 2 if ws_meas.max_row >= 2 else None
    perf_style_row = 2 if ws_perf.max_row >= 2 else None
    comp_style_row = 2 if ws_comp.max_row >= 2 else None

    perf_existing = extract_records(ws_perf, perf_headers)
    comp_existing = extract_records(ws_comp, comp_headers)

    pr_records = build_practitioner_role_records(pr_headers)
    meas_records = build_measure_records(meas_headers, measure_meta)
    perf_records = build_performance_records(perf_headers, perf_existing, period_results)
    comp_records = build_comparator_records(comp_headers, comp_existing, period_results)

    clear_data_rows(ws_pr)
    clear_data_rows(ws_meas)
    clear_data_rows(ws_perf)
    clear_data_rows(ws_comp)

    write_records(ws_pr, pr_headers, pr_records, pr_style_row)
    write_records(ws_meas, meas_headers, meas_records, meas_style_row)
    write_records(ws_perf, perf_headers, perf_records, perf_style_row)
    write_records(ws_comp, comp_headers, comp_records, comp_style_row)

    wb.save(output_path)

    return len(pr_records), len(meas_records), len(perf_records), len(comp_records)


# -----------------------------------------------------------------------------
# Main
# -----------------------------------------------------------------------------
def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Convert a Registry workbook into a populated SCAFFOLD ingestion workbook "
            "using a blank ingestion-model template."
        )
    )
    parser.add_argument("--registry", required=True, help="Path to the Registry workbook.")
    parser.add_argument("--template", required=True, help="Path to the blank SCAFFOLD ingestion-model template.")
    parser.add_argument("--output", required=True, help="Path where the populated workbook will be written.")
    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()

    try:
        registry_path = resolve_excel_path(args.registry, "Registry")
        scaffold_template_path = resolve_excel_path(args.template, "Template")
        output_path = Path(args.output).expanduser()
        output_path.parent.mkdir(parents=True, exist_ok=True)

        if output_path.resolve() == registry_path.resolve():
            raise ValueError("Output path must be different from the Registry workbook path.")
        if output_path.resolve() == scaffold_template_path.resolve():
            raise ValueError("Output path must be different from the template workbook path.")

    except Exception as exc:
        print(f"ERROR: {exc}")
        return 1

    all_warnings: List[str] = []
    mapping_info: Dict[str, str] = {}
    processing_stats: Dict[str, int] = {}
    assumptions: List[str] = []

    try:
        meetings, people, attendance, warnings_registry, mapping_info = load_registry_data(registry_path)
        all_warnings.extend(warnings_registry)

        measure_meta = default_measure_meta()
        assumptions.append("Updated template uses attendance-01 only; role groups are represented by subject identifiers.")

        period_results, warnings_calc, processing_stats = calculate_period_results(meetings, people, attendance)
        all_warnings.extend(warnings_calc)

        pr_count, meas_count, perf_count, comp_count = write_output_workbook(
            scaffold_template_path=scaffold_template_path,
            output_path=output_path,
            measure_meta=measure_meta,
            period_results=period_results,
        )

    except Exception as exc:
        print(f"ERROR: {exc}")
        return 1

    print("\nCompleted successfully.")
    print(f"Output: {output_path.resolve()}")
    print("Rows written:")
    print(f"  - PractitionerRole: {pr_count}")
    print(f"  - Measure: {meas_count}")
    print(f"  - PerformanceMeasureReport: {perf_count}")
    print(f"  - ComparatorMeasureReport: {comp_count}")

    print("Unprocessed records:")
    print(f"  - Attendance rows total: {processing_stats.get('attendance_rows_total', 0)}")
    print(f"  - Attendance rows valid: {processing_stats.get('attendance_rows_valid', 0)}")
    print(f"  - Attendance rows not processed: {processing_stats.get('attendance_rows_unprocessed', 0)}")

    print("\nValidation issues:")
    if all_warnings:
        for w in all_warnings:
            print(f"  - {w}")
    else:
        print("  - None")

    print("\nChanged worksheet mappings:")
    print("  - Registry.Meetings -> SCAFFOLD PerformanceMeasureReport/ComparatorMeasureReport period fields")
    print("  - Registry.People roles -> SCAFFOLD PerformanceMeasureReport/ComparatorMeasureReport grouped role rows")
    print("  - SCAFFOLD PractitionerRole -> role-group lookup rows (DLHS member/Faculty/Staff/Student)")
    print("  - Registry.Attendance Events -> SCAFFOLD PerformanceMeasureReport/ComparatorMeasureReport rate fields")

    print("Changed column mappings:")
    print(f"  - Registry attendance column accepted names: {ATTENDANCE_COLUMN_CANDIDATES}")
    if "registry.AttendanceEvents.Attendance" in mapping_info:
        print(f"  - Registry attendance column selected: {mapping_info['registry.AttendanceEvents.Attendance']}")

    print("Code conversions:")
    print("  - Role values normalized to Faculty/Staff/Student")
    print("  - Attendance values normalized to in_person/remote/not_attended")
    print("  - Subject identifier mapping: DLHS member=9999, Faculty=1119, Staff=1121, Student=1120")
    print("  - Comparator role codes now sourced from role-group rows (DLHS member/Faculty/Staff/Student)")

    print("Assumptions:")
    for assumption in assumptions:
        print(f"  - {assumption}")

    return 0


if __name__ == "__main__":
    sys.exit(main())
